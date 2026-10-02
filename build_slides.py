from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
OUT = ROOT / '03-slide-bai-giang-mang-may-tinh-osi.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = RGBColor(10, 20, 37)
PANEL = RGBColor(19, 34, 56)
PANEL2 = RGBColor(24, 45, 69)
WHITE = RGBColor(245, 248, 252)
MUTED = RGBColor(168, 186, 205)
CYAN = RGBColor(62, 211, 209)
BLUE = RGBColor(95, 155, 255)
AMBER = RGBColor(255, 193, 86)
GREEN = RGBColor(115, 222, 150)
FONT = 'Aptos'
MONO = 'Consolas'

def rect(s, x, y, w, h, fill=PANEL, line=None, radius=True):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background() if line is None else None
    if line is not None: sh.line.color.rgb = line
    return sh

def txt(s, text, x, y, w, h, size=20, color=WHITE, bold=False, font=FONT, align=None):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = font; p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = color; p.space_after = Pt(7)
        if align: p.alignment = align
    return sh

def base(title, tag='MẠNG MÁY TÍNH'):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    rect(s, 0, 0, .13, 7.5, CYAN, radius=False)
    txt(s, tag, .47, .28, 5, .28, 10, CYAN, True)
    txt(s, title, .47, .66, 12.25, .67, 27, WHITE, True)
    rect(s, .48, 7.11, 12.25, .014, PANEL2, radius=False)
    txt(s, 'OSI • TCP/IP • đường đi gói tin • NAT', .5, 7.17, 7, .2, 9, MUTED)
    txt(s, f'{len(prs.slides):02d}', 12.1, 7.15, .55, .24, 10, MUTED, align=PP_ALIGN.RIGHT)
    return s

def bullets(title, items, tag='MẠNG MÁY TÍNH', size=21):
    s=base(title,tag)
    for i,item in enumerate(items):
        y=1.55+i*(5.2/max(len(items),1))
        rect(s,.52,y,.1,.1,CYAN,radius=False)
        txt(s,item,.85,y-.11,11.8,5.0/max(len(items),1),size)
    return s

def cards(title, content, tag='MẠNG MÁY TÍNH', cols=2, fs=17):
    s=base(title,tag); rows=(len(content)+cols-1)//cols
    gap=.22; cw=(12.25-gap*(cols-1))/cols; ch=(5.42-gap*(rows-1))/rows
    for i,(head,body) in enumerate(content):
        c=i%cols; r=i//cols; x=.5+c*(cw+gap); y=1.5+r*(ch+gap)
        rect(s,x,y,cw,ch,PANEL)
        txt(s,head,x+.2,y+.18,cw-.4,.42,19,CYAN,True)
        txt(s,body,x+.2,y+.7,cw-.4,ch-.8,fs,WHITE)
    return s

def code(title, text, caption='', tag='GÓI TIN / DATA STRUCTURE', fs=20):
    s=base(title,tag)
    rect(s,.5,1.52,12.25,4.7,PANEL)
    txt(s,text,.82,1.8,11.65,4.2,fs,WHITE,False,MONO)
    if caption: txt(s,caption,.7,6.35,11.9,.47,15,AMBER)
    return s

def table(title, heads, rows, widths=None, fs=15, tag='MẠNG MÁY TÍNH'):
    s=base(title,tag); n=len(heads)
    widths=widths or [12.2/n]*n
    x0=.53; y0=1.53; rh=min(.68,5.3/(len(rows)+1))
    x=x0
    for j,h in enumerate(heads):
        rect(s,x,y0,widths[j]-.04,rh,PANEL2)
        txt(s,h,x+.11,y0+.13,widths[j]-.17,rh-.12,fs,CYAN,True)
        x+=widths[j]
    for i,row in enumerate(rows):
        x=x0; y=y0+(i+1)*rh
        for j,v in enumerate(row):
            rect(s,x,y,widths[j]-.04,rh-.035,PANEL if i%2==0 else RGBColor(17,30,49))
            txt(s,str(v),x+.11,y+.08,widths[j]-.19,rh-.08,fs,WHITE)
            x+=widths[j]
    return s

s=base('Mạng máy tính: từ socket đến dây mạng','BÀI GIẢNG 110 PHÚT')
txt(s,'7 tầng OSI • interface giữa các tầng • cấu trúc dữ liệu • NAT',.62,1.8,11.8,1.1,29,CYAN,True)
txt(s,'Dành cho người đã biết lập trình/socket cơ bản',.65,3.2,11.5,.5,21,WHITE)
rect(s,.65,4.35,11.9,1.47,PANEL)
txt(s,'Câu hỏi xuyên suốt',.9,4.58,3.7,.4,18,AMBER,True)
txt(s,'send() đi qua những tầng nào — và server trả lời về bằng cách nào?',.9,5.03,11.2,.53,22,WHITE)

bullets('Học xong, bạn sẽ làm được',[
    'Vẽ đường đi: ứng dụng → socket → TCP/UDP → IP → link → PHY.',
    'Đọc header Ethernet, IPv4, TCP/UDP và giải thích trường phân kênh.',
    'Theo dõi MAC, IP, port, TTL qua router và NAT.',
    'Thiết kế đường cloud → thiết bị sau NAT/CGNAT.'
],size=21)

table('Lịch buổi học 110 phút',['Thời gian','Nội dung','Đầu ra'],[
    ('0–20′','Thiết bị, OSI/TCP-IP, đóng gói','Vẽ PDU'),
    ('20–50′','Interface, header, cấu trúc stack','Đọc header'),
    ('50–79′','PHY/Link, IP, routing','Theo gói A→B'),
    ('79–93′','TCP/UDP, HTTPS, NAT/CGNAT','Giải thích reverse'),
    ('93–110′','Demo, bài tập, tổng kết','Thiết kế IoT')
],[1.7,6.1,4.4],16)

cards('Các thiết bị trong cùng một đường đi',[
    ('Host + NIC','Host chạy ứng dụng. NIC/driver nối host với Ethernet, Wi-Fi hoặc mạng khác.'),
    ('Switch + Access Point','Switch chuyển frame theo MAC trong LAN/VLAN. AP nối thiết bị Wi-Fi vào mạng.'),
    ('Router + gateway','Router chọn chặng tiếp theo theo IP. Default gateway là next hop mặc định của host.'),
    ('Modem/ONT + ISP','Modem/ONT kết thúc đường truyền. ISP cung cấp kết nối và địa chỉ phía ngoài.')
],fs=17)

table('7 tầng OSI: bản đồ nhanh',['Tầng','Tên','PDU','Ví dụ'],[
    ('7','Application','Data','HTTP, DNS, MQTT'),('6','Presentation','Data','UTF-8, JSON, TLS*'),
    ('5','Session','Data','RPC, phiên app'),('4','Transport','Segment/datagram','TCP, UDP, QUIC*'),
    ('3','Network','Packet','IPv4/IPv6, ICMP'),('2','Data Link','Frame','Ethernet, Wi-Fi, ARP*'),
    ('1','Physical','Bit','Đồng, quang, radio')
],[.8,3.05,2.35,6],15,tag='MÔ HÌNH OSI')

cards('OSI là mô hình; TCP/IP là stack sử dụng',[
    ('TCP/IP Application','Gộp nhiều chức năng OSI 5–7. TLS không nằm gọn trong một lớp.'),
    ('TCP/IP Transport','TCP/UDP; QUIC chạy trên UDP và cung cấp transport riêng.'),
    ('TCP/IP Internet','IP chuyển packet qua nhiều mạng; router tra địa chỉ IP đích.'),
    ('TCP/IP Link','Ethernet/Wi-Fi/PPP và PHY. Frame chỉ có ý nghĩa trên liên kết hiện tại.')
],tag='MÔ HÌNH OSI',fs=18)

code('Đóng gói và giải đóng gói',
     'HTTP data\n  ↓ + TCP header        → TCP segment\n  ↓ + IPv4 header       → IP packet\n  ↓ + Ethernet header/FCS → frame\n  ↓ tín hiệu            → bit trên dây/sóng',
     'Một lần send() không nhất thiết tương ứng một packet.',fs=22)

code('Interface ≠ trường phân kênh',
     'App  ── Socket API ──> TCP/UDP\nTCP/UDP ── IP Protocol/Next Header ──> IP\nIP ── EtherType (trên Ethernet) ──> Link\nLink ── MAC/PHY + driver ──> tín hiệu',
     'Socket API là interface; Protocol, EtherType, port là trường để phân luồng khi nhận.',fs=20)

table('Chiều nhận: tìm đúng bộ xử lý',['Trường','Giá trị ví dụ','Giao cho'],[
    ('EtherType','0x0800 / 0x86DD','IPv4 / IPv6'),
    ('IPv4 Protocol','6 / 17 / 1','TCP / UDP / ICMPv4'),
    ('TCP/UDP dst port','443 / 53 / 1883','Endpoint/socket phù hợp'),
    ('TCP 5-tuple','Protocol + 2 IP + 2 port','Connected socket cụ thể')
],[2.8,3.4,6],16,tag='INTERFACE GIỮA CÁC TẦNG')

code('Từ connect() và send() xuống NIC',
     'App:     socket() → connect() → send()\nTCP:     chọn port nguồn, state, seq, buffer\nIP:      chọn IP nguồn, route, next hop\nNeighbor: IP next hop → MAC\nLink:    thêm frame; driver → TX ring → NIC/PHY',
     'Offload của NIC có thể thay đổi thời điểm chia segment và tính checksum.',fs=19)

table('Header tối thiểu cần đọc',['Định dạng','Trường quan trọng','Kích thước'],[
    ('Ethernet II','dst MAC, src MAC, EtherType; FCS sau payload','14 B + 4 B FCS'),
    ('IPv4','src/dst IP, TTL, Protocol, checksum','20 B + options'),
    ('TCP','ports, seq, ack, flags, window, checksum','20 B + options'),
    ('UDP','ports, length, checksum','8 B'),
    ('ARP IPv4/Ethernet','opcode, sender/target MAC+IP','28 B thông thường')
],[2.4,7.4,2.4],15,tag='DATA STRUCTURE')

cards('Cấu trúc dữ liệu của stack',[
    ('Packet buffer','Linux sk_buff; lwIP pbuf. Giữ bytes, headroom và metadata.'),
    ('Socket / PCB / TCB','5-tuple, state TCP, seq/ack, window, timer, buffer.'),
    ('FIB + neighbor table','Prefix → next hop/interface; next-hop IP → MAC.'),
    ('Switch FDB + NAT table','MAC → cổng switch; mapping IP:port và trạng thái NAT.')
],tag='DATA STRUCTURE',fs=17)

cards('Ethernet, switch và ARP',[
    ('Frame Ethernet','MAC đích | MAC nguồn | EtherType | payload | FCS. VLAN tag thêm 4 byte.'),
    ('Switch','Học MAC nguồn → cổng; forward theo MAC đích; flood khi chưa biết.'),
    ('ARP','IPv4 next hop → MAC trên cùng liên kết. Khác subnet: ARP gateway.'),
    ('VLAN / STP','VLAN chia miền broadcast; STP giảm vòng lặp L2.')
],tag='TẦNG 2',fs=17)

table('Chuẩn Ethernet: chọn đúng media',['Chuẩn','Tốc độ','Môi trường'],[
    ('10/100/1000BASE-T','10/100/1000 Mb/s','Cáp xoắn đôi'),
    ('2.5/5/10GBASE-T','2.5/5/10 Gb/s','Cáp xoắn đôi phù hợp'),
    ('1000BASE-SX/LX','1 Gb/s','Quang đa/đơn mode'),
    ('100/1000BASE-T1','100 Mb/s / 1 Gb/s','Một đôi dây, automotive'),
    ('10BASE-T1L','10 Mb/s','Single Pair Ethernet đường dài')
],[4,3.15,5.05],15,tag='TẦNG 1–2')

cards('MAC, PHY và Wi-Fi',[
    ('Ethernet MCU','MAC trong MCU ↔ RMII/RGMII ↔ PHY; MDIO quản lý PHY.'),
    ('Wi-Fi chip','MAC + radio PHY; host kết nối qua SDIO/SPI/USB/PCIe.'),
    ('802.11','Scan → authenticate → associate; CSMA/CA; WPA2/WPA3.'),
    ('Wi-Fi 4/5/6/7','802.11n/ac/ax/be; tốc độ PHY lý thuyết ≠ tốc độ ứng dụng.')
],tag='TẦNG 1–2',fs=17)

cards('Các link khác trong embedded',[
    ('BLE / 802.15.4','BLE có stack riêng. Thread chạy IPv6/6LoWPAN trên 802.15.4.'),
    ('Zigbee / Z-Wave / LoRaWAN','Các mạng/giao thức công suất thấp; không tự động là Ethernet hay IP.'),
    ('Cellular / PPP','Modem 4G/5G hoặc LTE-M/NB-IoT có thể cấp đường IP; PPP là liên kết L2.'),
    ('CAN / RS-485 / UART','CAN có frame/bus; RS-485 là chuẩn điện; UART là giao diện nối tiếp.')
],tag='TẦNG 1–2',fs=16)

code('CIDR: gửi trực tiếp hay qua gateway?',
     'A: 192.168.1.10/24   mask 255.255.255.0\nB: 198.51.100.20\nA & mask = 192.168.1.0\nB & mask = 198.51.100.0\nKhác mạng → route mặc định qua 192.168.1.1',
     'Máy có nhiều route chọn prefix khớp dài nhất.',tag='TẦNG 3',fs=20)

code('Router làm gì ở mỗi chặng?',
     'Nhận frame → lấy IP packet\nĐọc IP đích → tra FIB (longest prefix match)\nGiảm TTL IPv4 / Hop Limit IPv6\nCập nhật IPv4 header checksum\nTìm link next hop → đóng frame mới → gửi',
     'Switch L2 không giảm TTL; NAT là chức năng bổ sung.',tag='TẦNG 3',fs=20)

table('Theo dõi gói A → server qua NAT',['Vị trí','MAC đích','IP:port nguồn','IP:port đích','TTL'],[
    ('LAN A','Gateway LAN','192.168.1.10:51514','198.51.100.20:443','64'),
    ('Sau NAT','Next hop WAN','203.0.113.10:62001','198.51.100.20:443','63'),
    ('Gần server','Server B','203.0.113.10:62001','198.51.100.20:443','<63')
],[1.45,2.18,3.5,3.5,1.57],13,tag='HÀNH TRÌNH GÓI TIN')

cards('TTL, ICMP và MTU',[
    ('TTL / Hop Limit','Mỗi router giảm 1; hết giá trị thì bỏ packet. Traceroute tận dụng ICMP Time Exceeded.'),
    ('Path MTU','Link có MTU nhỏ nhất giới hạn packet; IPv4 DF + ICMP báo cần giảm kích thước.'),
    ('IPv6','Router không phân mảnh; gửi ICMPv6 Packet Too Big cho nguồn.'),
    ('Chẩn đoán','Một hop traceroute im lặng có thể do chặn probe; chưa đủ kết luận đường đứt.')
],tag='TẦNG 3',fs=16)

table('Transport: TCP, UDP, QUIC',['Thuộc tính','TCP','UDP','QUIC'],[
    ('Dữ liệu','Byte stream','Datagram','Streams trên UDP'),
    ('Tin cậy/thứ tự','Do TCP cung cấp','Ứng dụng tự xử lý','QUIC cung cấp cho stream'),
    ('Bắt tay','SYN/SYN-ACK/ACK','Không có','Tích hợp TLS'),
    ('Ví dụ','SSH, MQTT, HTTP/2','DNS, DHCP, CoAP','HTTP/3')
],[2.35,3.1,3.1,3.65],14,tag='TẦNG 4')

cards('Port, socket và 5-tuple',[
    ('Port','Định danh transport endpoint. IANA: 0–1023, 1024–49151, 49152–65535.'),
    ('TCP connection','(TCP, src IP, src port, dst IP, dst port) phân biệt các kết nối.'),
    ('Listening socket','bind + listen ở server; accept tạo socket kết nối riêng.'),
    ('Ephemeral port','OS chọn port nguồn tạm thời; vùng thực tế có thể được cấu hình khác IANA.')
],tag='TẦNG 4',fs=17)

cards('Ứng dụng, dữ liệu và bảo mật',[
    ('HTTP','Request/response, method, status, header, cookie. HTTP/2 multiplex trên TCP.'),
    ('TLS','Xác thực certificate, thỏa thuận khóa và bảo vệ dữ liệu. HTTPS/TCP: TLS giữa HTTP và TCP.'),
    ('DNS / DHCP','DNS phân giải tên; DHCP cấp cấu hình IP/GW/DNS trong mạng.'),
    ('IoT','MQTT qua TCP; CoAP thường qua UDP; HTTP/3 chạy trên QUIC/UDP.')
],tag='TẦNG 5–7',fs=17)

code('Gõ https://example.com rồi Enter',
     'DNS tên → IP (hoặc cache)\nRoute → next hop; ARP tìm MAC gateway\nTCP handshake → TLS handshake\nHTTP request được gửi trong TLS\nTCP/IP/Link đóng gói → router/NAT → server\nServer giải đóng gói và trả response',
     'Đây là ví dụ HTTP/2 trên TCP/IPv4; HTTP/3 dùng QUIC/UDP.',tag='VÍ DỤ XUYÊN SUỐT',fs=19)

cards('NAT/PAT và IP public',[
    ('Bên trong','192.168.1.10:51514 → 198.51.100.20:443'),
    ('Bản ghi NAT','192.168.1.10:51514 ↔ 203.0.113.10:62001; gắn trạng thái luồng.'),
    ('Server nhìn thấy','203.0.113.10:62001, không thấy trực tiếp IP LAN của client.'),
    ('CGNAT','ISP chia một IP public cho nhiều thuê bao; WAN 100.64.0.0/10 là dấu hiệu thường gặp.')
],tag='NAT / IP PUBLIC',fs=17)

cards('Server có kết nối ngược được không?',[
    ('Cùng kết nối: CÓ','TCP full duplex. Server gửi dữ liệu qua socket client đã mở; NAT có trạng thái.'),
    ('Kết nối TCP mới: thường KHÔNG','Cần service listen, route/NAT rule và firewall cho phép.'),
    ('Port nguồn cũ','Ephemeral port trên client không tự thành cổng lắng nghe SYN mới.'),
    ('Sau CGNAT','Người dùng thường không cấu hình được NAT của ISP; IP public có thể dùng chung.')
],tag='NAT / KẾT NỐI NGƯỢC',fs=17)

table('Cách đưa lệnh cloud → thiết bị',['Cách','Cơ chế','Điều kiện'],[
    ('MQTT/WebSocket','Client mở và giữ kết nối','Heartbeat, reconnect, xác thực'),
    ('Polling','Client hỏi việc định kỳ','Chấp nhận độ trễ'),
    ('Port forwarding','Chuyển WAN port → LAN host','Inbound path + firewall'),
    ('IPv6','Địa chỉ định tuyến toàn cầu','Firewall cho phép'),
    ('Tunnel / VPN','Client mở tunnel ra ngoài','Điểm trung gian'),
    ('ICE / TURN','Thử P2P hoặc relay','NAT phù hợp / relay')
],[2.8,4.6,4.8],14,tag='NAT / IOT')

table('Demo trực tiếp',['Mục tiêu','Linux','Windows'],[
    ('IP cục bộ','ip addr','ipconfig /all'),
    ('Bảng route','ip route','route print'),
    ('Neighbor/ARP','ip neigh / arp -a','arp -a'),
    ('Các hop','traceroute example.com','tracert example.com'),
    ('IP bên ngoài','curl -4 https://ifconfig.me','curl.exe -4 https://ifconfig.me')
],[2.55,4.8,4.85],15,tag='THỰC HÀNH')

bullets('Bài tập 5 phút',[
    'EtherType 0x0800 → IP Protocol 17 → UDP dst port 53: ứng dụng nào có thể là?',
    '192.168.10.17/26 gửi đến .50 và .90: địa chỉ nào cùng subnet?',
    'NAT đổi 192.168.1.10:51514 thành 203.0.113.10:62001: server thấy gì?',
    'Thiết kế cloud gửi lệnh cho thiết bị sau CGNAT, có reconnect và chống lệnh trùng.'
],tag='BÀI TẬP',size=18)

cards('4 điều cần nhớ',[
    ('01 — Ranh giới','Socket API là interface lập trình; EtherType/Protocol/port dùng phân kênh.'),
    ('02 — Địa chỉ','MAC theo liên kết; IP qua mạng; port tìm transport endpoint.'),
    ('03 — Router + NAT','Router tạo frame mới và giảm TTL; NAT có thể đổi IP/port.'),
    ('04 — Cloud → IoT','Dùng kết nối chủ động ra ngoài; duy trì, reconnect, ACK lệnh.')
],tag='TỔNG KẾT',fs=17)

s=base('Tài liệu chuẩn và giáo án đầy đủ','TÀI LIỆU')
txt(s,'01-outline-mang-may-tinh-osi.md\n02-bai-giang-chi-tiet-mang-may-tinh.md',.7,1.6,12,1.1,22,CYAN,True,MONO)
txt(s,'RFC 791 · 826 · 1918 · 3022 · 4787 · 6598 · 8200 · 8445 · 9000 · 9114 · 9293',.72,3.2,11.7,1.1,21,WHITE)
txt(s,'IANA Service Name and Transport Protocol Port Number Registry',.72,4.65,11.7,.7,20,MUTED)
txt(s,'Các IP 198.51.100.0/24 và 203.0.113.0/24 trong slide chỉ dùng làm ví dụ.',.72,5.75,11.7,.6,17,AMBER)

prs.save(OUT)
print(f'{OUT}\nSlides: {len(prs.slides)}')
